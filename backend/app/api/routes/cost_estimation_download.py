from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import Response
from sqlalchemy.orm import Session
from typing import List
from decimal import Decimal
from pydantic import BaseModel

from app.database.session import get_db
from app.core.dependencies import get_current_user
from app.core.rbac import require_admin_or_user
from app.services.cost_service import CostCalculationService
from app.models.user import User

router = APIRouter(prefix="/api/v1/cost-estimation", tags=["Cost Estimation Download"])


class AdditionalCostItem(BaseModel):
    cost_name: str
    cost_value: float


class CostDownloadRequest(BaseModel):
    customer_id: int
    product_id: int
    quantity: float
    additional_costs: List[AdditionalCostItem] = []


@router.post("/download")
def download_cost_sheet(
    request: CostDownloadRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin_or_user)
):
    """
    Generate and download a detailed cost sheet as Excel file.
    Uses the same calculation service as the calculate endpoint.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from io import BytesIO
    from datetime import datetime
    
    # Validate customer exists
    from app.models.customer import Customer
    customer = db.query(Customer).filter(Customer.id == request.customer_id).first()
    if not customer:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Customer with id {request.customer_id} not found"
        )
    
    # Validate product exists
    from app.models.product import Product
    product = db.query(Product).filter(Product.id == request.product_id).first()
    if not product:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Product with id {request.product_id} not found"
        )
    
    # Calculate cost using the same service
    from app.models.part import Part
    from app.models.assembly import Assembly
    from app.models.operation import Operation
    from app.models.machine import Machine
    
    parts = db.query(Part).filter(Part.product_id == request.product_id).all()
    assemblies = db.query(Assembly).filter(Assembly.product_id == request.product_id).all()
    
    part_ids = [p.id for p in parts]
    operations_by_part = {}
    if part_ids:
        operations = db.query(Operation).filter(Operation.part_id.in_(part_ids)).all()
        for op in operations:
            if op.part_id not in operations_by_part:
                operations_by_part[op.part_id] = []
            operations_by_part[op.part_id].append(op)
    
    machine_ids = set()
    for ops in operations_by_part.values():
        for op in ops:
            machine_ids.add(op.machine_id)
    
    machines_map = {}
    if machine_ids:
        machines = db.query(Machine).filter(Machine.id.in_(machine_ids)).all()
        machines_map = {m.id: m for m in machines}
    
    quantity = Decimal(str(request.quantity))
    cost_breakdown = CostCalculationService.calculate_product_cost(
        product, parts, assemblies, operations_by_part, machines_map, quantity
    )
    
    # Calculate additional costs total
    additional_costs_total = Decimal("0")
    for cost_item in request.additional_costs:
        additional_costs_total += Decimal(str(cost_item.cost_value))
    
    # Calculate grand total
    grand_total = Decimal(str(cost_breakdown["total_manufacturing_cost"])) + additional_costs_total
    
    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cost Estimation"
    
    # Define styles
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write header information
    ws['A1'] = "COST ESTIMATION SHEET"
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal="center")
    ws.merge_cells('A1:E1')
    
    ws['A3'] = "Customer:"
    ws['B3'] = customer.company_name
    ws['A4'] = "Product:"
    ws['B4'] = f"{product.product_number} - {product.product_name}"
    ws['A5'] = "Quantity:"
    ws['B5'] = float(quantity)
    ws['A6'] = "Date:"
    ws['B6'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    # Write cost breakdown
    row = 8
    ws[f'A{row}'] = "Type"
    ws[f'B{row}'] = "Number/Name"
    ws[f'C{row}'] = "Description"
    ws[f'D{row}'] = "Quantity"
    ws[f'E{row}'] = "Cost"
    
    for col in ['A', 'B', 'C', 'D', 'E']:
        cell = ws[f'{col}{row}']
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    row += 1
    
    # Write direct parts
    for part_cost in cost_breakdown["direct_part_costs"]:
        ws[f'A{row}'] = "Part"
        ws[f'B{row}'] = f"{part_cost['part_number']} - {part_cost['part_name']}"
        ws[f'C{row}'] = part_cost.get('description', '')
        ws[f'D{row}'] = part_cost['quantity']
        ws[f'E{row}'] = part_cost['total_part_cost']
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            ws[f'{col}{row}'].border = thin_border
        
        row += 1
    
    # Write assemblies (simplified for now)
    for assembly_cost in cost_breakdown["assembly_costs"]:
        ws[f'A{row}'] = "Assembly"
        ws[f'B{row}'] = f"{assembly_cost['assembly_number']} - {assembly_cost['assembly_name']}"
        ws[f'C{row}'] = ""
        ws[f'D{row}'] = assembly_cost['quantity']
        ws[f'E{row}'] = assembly_cost['total_assembly_cost']
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            ws[f'{col}{row}'].border = thin_border
        
        row += 1
    
    # Write manufacturing total
    row += 1
    ws[f'D{row}'] = "Manufacturing Total:"
    ws[f'E{row}'] = float(cost_breakdown["total_manufacturing_cost"])
    ws[f'D{row}'].font = header_font
    ws[f'E{row}'].font = header_font
    
    # Write additional costs
    row += 2
    ws[f'A{row}'] = "Additional Costs"
    ws[f'A{row}'].font = header_font
    row += 1
    
    for cost_item in request.additional_costs:
        ws[f'A{row}'] = cost_item.cost_name
        ws[f'E{row}'] = cost_item.cost_value
        ws[f'A{row}'].border = thin_border
        ws[f'E{row}'].border = thin_border
        row += 1
    
    # Write additional costs total
    ws[f'D{row}'] = "Additional Costs Total:"
    ws[f'E{row}'] = float(additional_costs_total)
    ws[f'D{row}'].font = header_font
    ws[f'E{row}'].font = header_font
    
    # Write grand total
    row += 1
    ws[f'D{row}'] = "GRAND TOTAL:"
    ws[f'E{row}'] = float(grand_total)
    ws[f'D{row}'].font = Font(bold=True, size=14)
    ws[f'E{row}'].font = Font(bold=True, size=14)
    ws[f'D{row}'].fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    ws[f'E{row}'].fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    
    # Auto-adjust column widths
    for col in ['A', 'B', 'C', 'D', 'E']:
        max_length = 0
        for cell in ws[col]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col].width = max_length + 2
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Cost_Estimation_{product.product_number}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
